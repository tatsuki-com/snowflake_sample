"""openpyxl 共通操作モジュール"""

import os

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string


def open_workbook(filepath: str, read_only: bool = False, keep_vba: bool = False):
    """ワークブックを開いて返す。"""
    return load_workbook(filepath, read_only=read_only, keep_vba=keep_vba)


def get_or_create_workbook(filepath: str, keep_vba: bool = False) -> Workbook:
    """ファイルが存在すれば開き、なければ新規作成して返す。"""
    if os.path.exists(filepath):
        return load_workbook(filepath, keep_vba=keep_vba)
    return Workbook()


def get_sheet(workbook: Workbook, sheet_name: str):
    """指定名のシートを返す。存在しなければ None を返す。"""
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return None


def get_or_create_sheet(workbook: Workbook, sheet_name: str):
    """指定名のシートを返す。存在しなければ新規作成して返す。"""
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return workbook.create_sheet(sheet_name)


def get_next_row(sheet) -> int:
    """シートの最終行の次の行番号を返す。"""
    return sheet.max_row + 1


def col_letter_to_index(letter: str) -> int:
    """列文字（A, B, C...）を1始まりのインデックスに変換する。"""
    return column_index_from_string(letter)


def write_cell(sheet, row: int, col_letter: str, value):
    """指定セルに値を書き込む。"""
    col_idx = col_letter_to_index(col_letter)
    sheet.cell(row=row, column=col_idx, value=value)


def read_cell(sheet, cell_ref: str):
    """セル参照（例: A1, B11）から値を読み取る。"""
    return sheet[cell_ref].value
