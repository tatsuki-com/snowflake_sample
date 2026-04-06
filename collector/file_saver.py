"""添付ファイル保存・重複検知・リネームモジュール"""

import os
import logging
from datetime import datetime

from openpyxl import load_workbook

from common.excel_utils import get_or_create_sheet

logger = logging.getLogger("email_attachment_system")


def save_attachment(attachment, folder_a: str, received_time: datetime,
                    datetime_sheet_name: str) -> dict:
    """添付ファイルをフォルダAに保存し、_get_datetimeシートを追加する。

    Returns:
        dict: 保存結果情報
            - filename: 保存後のファイル名
            - original_filename: 元のファイル名
            - received_datetime: メール受信日時
            - collected_datetime: 収集処理日時
            - is_duplicate: 重複によりリネームされたか
            - remarks: 備考メッセージ
    """
    os.makedirs(folder_a, exist_ok=True)

    original_filename = attachment.FileName
    filepath = os.path.join(folder_a, original_filename)
    is_duplicate = False
    remarks = ""
    saved_filename = original_filename

    if os.path.exists(filepath):
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
        name, ext = os.path.splitext(original_filename)
        saved_filename = f"{name}_{timestamp}{ext}"
        filepath = os.path.join(folder_a, saved_filename)
        is_duplicate = True
        remarks = f"同名ファイル重複のためリネーム: {original_filename} -> {saved_filename}"
        logger.warning(remarks)

    try:
        attachment.SaveAsFile(filepath)
        collected_datetime = datetime.now()
        logger.info(f"添付ファイル保存: {saved_filename}")
    except Exception as e:
        logger.error(f"添付ファイル保存失敗: {original_filename} - {e}")
        raise

    try:
        _add_datetime_sheet(filepath, received_time, datetime_sheet_name)
    except Exception as e:
        logger.error(f"_get_datetimeシート追加失敗: {saved_filename} - {e}")
        raise

    return {
        "filename": saved_filename,
        "original_filename": original_filename,
        "received_datetime": received_time,
        "collected_datetime": collected_datetime,
        "is_duplicate": is_duplicate,
        "remarks": remarks,
    }


def _add_datetime_sheet(filepath: str, received_time: datetime,
                        sheet_name: str) -> None:
    """保存済みExcelファイルに_get_datetimeシートを追加し、A1に受信日時を記録する。"""
    wb = load_workbook(filepath)
    ws = get_or_create_sheet(wb, sheet_name)
    ws["A1"] = received_time
    wb.save(filepath)
    wb.close()
