"""file_writer のテスト"""

import os
import pytest
from datetime import datetime

from openpyxl import load_workbook

from aggregator.file_writer import write_to_aggregation_file, HEADER_ROW


def _make_config(tmp_path):
    folder_b = str(tmp_path / "aggregated")
    return {
        "paths": {"folder_b": folder_b},
        "files": {
            "aggregation_targets": {
                "type_1": "alpha.xlsx",
                "type_2": "beta.xlsx",
            },
        },
    }


class TestWriteToAggregationFile:
    def test_creates_new_file_with_header(self, tmp_path):
        config = _make_config(tmp_path)
        data = {
            "routing_type": "type_1",
            "received_datetime": datetime(2026, 4, 1, 9, 0, 0),
            "data_values": [f"val_{i}" for i in range(11)],
        }

        result = write_to_aggregation_file(data, config)
        assert result is True

        filepath = os.path.join(config["paths"]["folder_b"], "alpha.xlsx")
        wb = load_workbook(filepath)
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "受信日時"
        assert ws.cell(row=2, column=1).value == datetime(2026, 4, 1, 9, 0, 0)
        assert ws.cell(row=2, column=2).value == "val_0"
        wb.close()

    def test_appends_to_existing_file(self, tmp_path):
        config = _make_config(tmp_path)
        data1 = {
            "routing_type": "type_1",
            "received_datetime": datetime(2026, 4, 1, 9, 0, 0),
            "data_values": ["a"] * 11,
        }
        data2 = {
            "routing_type": "type_1",
            "received_datetime": datetime(2026, 4, 2, 10, 0, 0),
            "data_values": ["b"] * 11,
        }

        write_to_aggregation_file(data1, config)
        write_to_aggregation_file(data2, config)

        filepath = os.path.join(config["paths"]["folder_b"], "alpha.xlsx")
        wb = load_workbook(filepath)
        ws = wb.active
        assert ws.max_row == 3  # header + 2 data rows
        assert ws.cell(row=3, column=2).value == "b"
        wb.close()

    def test_unknown_routing_type(self, tmp_path):
        config = _make_config(tmp_path)
        data = {
            "routing_type": "unknown_type",
            "received_datetime": datetime(2026, 4, 1, 9, 0, 0),
            "data_values": [],
        }

        result = write_to_aggregation_file(data, config)
        assert result is False
