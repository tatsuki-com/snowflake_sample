"""file_saver のテスト (Outlook COM なしでテスト可能な部分)"""

import os
import pytest
from datetime import datetime
from unittest.mock import MagicMock

from openpyxl import Workbook

from collector.file_saver import save_attachment, _add_datetime_sheet


class TestAddDatetimeSheet:
    def test_adds_sheet_with_datetime(self, tmp_path):
        filepath = str(tmp_path / "test.xlsx")
        wb = Workbook()
        wb.save(filepath)
        wb.close()

        received_time = datetime(2026, 4, 1, 9, 0, 0)
        _add_datetime_sheet(filepath, received_time, "_get_datetime")

        from openpyxl import load_workbook
        wb = load_workbook(filepath)
        assert "_get_datetime" in wb.sheetnames
        ws = wb["_get_datetime"]
        assert ws["A1"].value == received_time
        wb.close()

    def test_overwrites_existing_sheet(self, tmp_path):
        filepath = str(tmp_path / "test.xlsx")
        wb = Workbook()
        ws = wb.create_sheet("_get_datetime")
        ws["A1"] = "old_value"
        wb.save(filepath)
        wb.close()

        new_time = datetime(2026, 4, 2, 10, 30, 0)
        _add_datetime_sheet(filepath, new_time, "_get_datetime")

        from openpyxl import load_workbook
        wb = load_workbook(filepath)
        assert wb["_get_datetime"]["A1"].value == new_time
        wb.close()


class TestSaveAttachment:
    def test_save_new_file(self, tmp_path):
        folder_a = str(tmp_path / "collected")

        # Excelファイルを事前に作成（SaveAsFileのモック用）
        mock_filepath = str(tmp_path / "source.xlsx")
        wb = Workbook()
        wb.save(mock_filepath)
        wb.close()

        attachment = MagicMock()
        attachment.FileName = "report.xlsx"

        def mock_save_as_file(path):
            import shutil
            shutil.copy2(mock_filepath, path)

        attachment.SaveAsFile = mock_save_as_file

        received_time = datetime(2026, 4, 1, 9, 0, 0)
        result = save_attachment(attachment, folder_a, received_time, "_get_datetime")

        assert result["filename"] == "report.xlsx"
        assert result["original_filename"] == "report.xlsx"
        assert result["is_duplicate"] is False
        assert result["remarks"] == ""
        assert os.path.exists(os.path.join(folder_a, "report.xlsx"))

    def test_save_duplicate_file(self, tmp_path):
        folder_a = str(tmp_path / "collected")
        os.makedirs(folder_a)

        # 既存ファイルを作成
        existing = os.path.join(folder_a, "report.xlsx")
        wb = Workbook()
        wb.save(existing)
        wb.close()

        mock_filepath = str(tmp_path / "source.xlsx")
        wb = Workbook()
        wb.save(mock_filepath)
        wb.close()

        attachment = MagicMock()
        attachment.FileName = "report.xlsx"

        def mock_save_as_file(path):
            import shutil
            shutil.copy2(mock_filepath, path)

        attachment.SaveAsFile = mock_save_as_file

        received_time = datetime(2026, 4, 1, 9, 0, 0)
        result = save_attachment(attachment, folder_a, received_time, "_get_datetime")

        assert result["is_duplicate"] is True
        assert result["filename"] != "report.xlsx"
        assert result["filename"].startswith("report_")
        assert result["original_filename"] == "report.xlsx"
