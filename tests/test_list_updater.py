"""list_updater のテスト"""

import os
import pytest
from datetime import datetime

from openpyxl import Workbook

from collector.list_updater import append_to_file_list, update_flags


def _make_config(tmp_path):
    file_list_path = str(tmp_path / "filelist.xlsx")
    return {
        "files": {"file_list_path": file_list_path},
        "file_list_sheet": {
            "name": "一覧",
            "columns": {
                "filename": "A",
                "original_filename": "B",
                "received_datetime": "C",
                "collected_datetime": "D",
                "duplicate_flag": "E",
                "updated_flag": "F",
                "routing_type": "G",
                "remarks": "H",
            },
        },
    }


class TestAppendToFileList:
    def test_append_creates_file_and_adds_row(self, tmp_path):
        config = _make_config(tmp_path)
        file_info = {
            "filename": "test.xlsx",
            "original_filename": "test.xlsx",
            "received_datetime": datetime(2026, 4, 1, 9, 0, 0),
            "collected_datetime": datetime(2026, 4, 1, 9, 5, 0),
            "is_duplicate": False,
            "remarks": "",
        }

        append_to_file_list(file_info, config)

        from openpyxl import load_workbook
        wb = load_workbook(config["files"]["file_list_path"])
        ws = wb["一覧"]
        # 新規ワークブックの場合 max_row=1 のため、データは2行目に書き込まれる
        assert ws["A2"].value == "test.xlsx"
        assert ws["E2"].value == 0
        assert ws["F2"].value == 0
        wb.close()

    def test_append_duplicate_flag(self, tmp_path):
        config = _make_config(tmp_path)
        file_info = {
            "filename": "test_20260401.xlsx",
            "original_filename": "test.xlsx",
            "received_datetime": datetime(2026, 4, 1, 9, 0, 0),
            "collected_datetime": datetime(2026, 4, 1, 9, 5, 0),
            "is_duplicate": True,
            "remarks": "重複リネーム",
        }

        append_to_file_list(file_info, config)

        from openpyxl import load_workbook
        wb = load_workbook(config["files"]["file_list_path"])
        ws = wb["一覧"]
        assert ws["A2"].value == "test_20260401.xlsx"
        assert ws["E2"].value == 1
        assert ws["H2"].value == "重複リネーム"
        wb.close()


class TestUpdateFlags:
    def test_updates_correct_row(self, tmp_path):
        config = _make_config(tmp_path)

        # 事前にファイル一覧を作成
        file_list_path = config["files"]["file_list_path"]
        wb = Workbook()
        ws = wb.active
        ws.title = "一覧"
        ws["A1"] = "file1.xlsx"
        ws["F1"] = 0
        ws["G1"] = ""
        ws["A2"] = "file2.xlsx"
        ws["F2"] = 0
        ws["G2"] = ""
        wb.save(file_list_path)
        wb.close()

        update_flags("file2.xlsx", "type_1", config)

        from openpyxl import load_workbook
        wb = load_workbook(config["files"]["file_list_path"])
        ws = wb["一覧"]
        assert ws["F1"].value == 0  # file1は変更されない
        assert ws["F2"].value == 1
        assert ws["G2"].value == "type_1"
        wb.close()
