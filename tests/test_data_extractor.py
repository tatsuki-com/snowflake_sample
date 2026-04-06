"""data_extractor のテスト"""

import os
import pytest
from datetime import datetime

from openpyxl import Workbook

from aggregator.data_extractor import extract_data


def _make_config():
    return {
        "data_extraction": {
            "datetime_sheet": "_get_datetime",
            "datetime_cell": "A1",
            "data_range": {"column": "B", "start_row": 10, "end_row": 20},
            "routing_cell": "B11",
        },
    }


def _create_test_file(filepath, routing_type="type_1", received_time=None, add_dt_sheet=True):
    if received_time is None:
        received_time = datetime(2026, 4, 1, 9, 0, 0)

    wb = Workbook()
    ws = wb.active
    ws.title = "メイン"
    ws["B11"] = routing_type

    if add_dt_sheet:
        dt_ws = wb.create_sheet("_get_datetime")
        dt_ws["A1"] = received_time
        for row in range(10, 21):
            dt_ws.cell(row=row, column=2, value=f"data_{row}")

    wb.save(filepath)
    wb.close()


class TestExtractData:
    def test_extract_valid_data(self, tmp_path):
        filepath = str(tmp_path / "test.xlsx")
        _create_test_file(filepath)

        config = _make_config()
        result = extract_data(filepath, config)

        assert result is not None
        assert result["routing_type"] == "type_1"
        assert result["received_datetime"] == datetime(2026, 4, 1, 9, 0, 0)
        assert len(result["data_values"]) == 11  # B10〜B20
        assert result["data_values"][0] == "data_10"

    def test_missing_datetime_sheet(self, tmp_path):
        filepath = str(tmp_path / "test.xlsx")
        _create_test_file(filepath, add_dt_sheet=False)

        config = _make_config()
        result = extract_data(filepath, config)
        assert result is None

    def test_empty_routing_cell(self, tmp_path):
        filepath = str(tmp_path / "test.xlsx")
        _create_test_file(filepath, routing_type="")

        config = _make_config()
        result = extract_data(filepath, config)
        assert result is None

    def test_none_routing_cell(self, tmp_path):
        filepath = str(tmp_path / "test.xlsx")
        _create_test_file(filepath, routing_type=None)

        config = _make_config()
        result = extract_data(filepath, config)
        assert result is None

    def test_file_not_found(self):
        config = _make_config()
        result = extract_data("/nonexistent/file.xlsx", config)
        assert result is None
