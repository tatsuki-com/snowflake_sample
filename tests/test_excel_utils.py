"""excel_utils のテスト"""

import os
import pytest
from openpyxl import Workbook

from common.excel_utils import (
    open_workbook,
    get_or_create_workbook,
    get_sheet,
    get_or_create_sheet,
    get_next_row,
    col_letter_to_index,
    write_cell,
    read_cell,
)


@pytest.fixture
def sample_workbook(tmp_path):
    filepath = tmp_path / "test.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "header"
    ws["B2"] = 42
    wb.save(str(filepath))
    wb.close()
    return str(filepath)


class TestOpenWorkbook:
    def test_open_existing(self, sample_workbook):
        wb = open_workbook(sample_workbook)
        assert "Sheet1" in wb.sheetnames
        wb.close()

    def test_open_nonexistent(self):
        with pytest.raises(FileNotFoundError):
            open_workbook("/nonexistent/file.xlsx")


class TestGetOrCreateWorkbook:
    def test_open_existing(self, sample_workbook):
        wb = get_or_create_workbook(sample_workbook)
        assert "Sheet1" in wb.sheetnames
        wb.close()

    def test_create_new(self, tmp_path):
        filepath = str(tmp_path / "new.xlsx")
        wb = get_or_create_workbook(filepath)
        assert wb is not None
        wb.close()


class TestGetSheet:
    def test_existing_sheet(self, sample_workbook):
        wb = open_workbook(sample_workbook)
        ws = get_sheet(wb, "Sheet1")
        assert ws is not None
        assert ws.title == "Sheet1"
        wb.close()

    def test_nonexistent_sheet(self, sample_workbook):
        wb = open_workbook(sample_workbook)
        ws = get_sheet(wb, "NonExistent")
        assert ws is None
        wb.close()


class TestGetOrCreateSheet:
    def test_existing_sheet(self, sample_workbook):
        wb = open_workbook(sample_workbook)
        ws = get_or_create_sheet(wb, "Sheet1")
        assert ws["A1"].value == "header"
        wb.close()

    def test_create_new_sheet(self, sample_workbook):
        wb = open_workbook(sample_workbook)
        ws = get_or_create_sheet(wb, "NewSheet")
        assert ws.title == "NewSheet"
        wb.close()


class TestGetNextRow:
    def test_with_data(self, sample_workbook):
        wb = open_workbook(sample_workbook)
        ws = wb["Sheet1"]
        assert get_next_row(ws) == 3  # row 1: header, row 2: data, next = 3
        wb.close()


class TestColLetterToIndex:
    def test_single_letter(self):
        assert col_letter_to_index("A") == 1
        assert col_letter_to_index("B") == 2
        assert col_letter_to_index("Z") == 26

    def test_double_letter(self):
        assert col_letter_to_index("AA") == 27


class TestWriteAndReadCell:
    def test_write_and_read(self, sample_workbook):
        wb = open_workbook(sample_workbook)
        ws = wb["Sheet1"]
        write_cell(ws, 3, "C", "test_value")
        assert read_cell(ws, "C3") == "test_value"
        wb.close()
